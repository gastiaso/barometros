"""Genera inventario/revision.xlsx: las tablas del inventario y de la revisión, con formato (los CSV son la versión canónica).
Uso: python scripts/06_revision_xlsx.py <raíz del proyecto>"""
import os, sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = sys.argv[1] if len(sys.argv) > 1 else "."
INV = os.path.join(ROOT, "inventario")
OUT = os.path.join(INV, "revision.xlsx")

leyenda = pd.DataFrame({"Campo": [
    "Objetivo", "Hojas", "Niveles (nivel_final)", "  A", "  B", "  C", "  AUSENTE", "revisado", "Comentarios",
    "Fuente",
], "Descripción": [
    "Revisión del crosswalk contra cuestionarios, libros de códigos y fichas técnicas (scripts/04_enunciados.py y 05_revision.py). Los CSV de inventario/ son la versión canónica.",
    "Crosswalk (revisado), Enunciados (versiones de cada pregunta), Rupturas, Decisiones (manuales), Partidos, Problemas, Candidatos, Códigos (cambios de significado), Estudios, Líderes.",
    "",
    "Comparable tal cual (puede cambiar el nombre de la variable o el texto de las etiquetas).",
    "Comparable con una regla de armonización documentada (cambio de redacción, códigos de no respuesta, códigos reutilizados en listas abiertas).",
    "No comparable: cambian las categorías, la definición o el contexto de la pregunta.",
    "El concepto no está en el estudio (ver nota: confirmado o derivable).",
    "manual = decisión razonada en decisiones_manuales.csv; auto = regla aplicada con cuestionario y .sav o solo con .sav.",
    "Columna «comentario» (amarillo) de la hoja Crosswalk, para anotaciones propias.",
    "Fuente de los datos: Centro de Investigaciones Sociológicas (microdatos y documentación de los barómetros, ene-2021 a oct-2026).",
]})

def csv_(n):
    return pd.read_csv(os.path.join(INV, n), dtype=str).fillna("")

cross = csv_("crosswalk_revisado.csv")
cross["comentario"] = ""
sheets = {
    "Leyenda": leyenda,
    "Crosswalk": cross,
    "Enunciados": csv_("enunciados_versiones.csv"),
    "Rupturas": csv_("rupturas_candidatas.csv"),
    "Decisiones": csv_("decisiones_manuales.csv"),
    "Partidos": csv_("partidos_armonizados.csv"),
    "Problemas": csv_("mapa_problemas.csv").drop_duplicates(["id_armonizado", "etiqueta"]),
    "Candidatos": csv_("mapa_candidatos.csv").drop_duplicates(["id_armonizado", "etiqueta"]),
    "Codigos": csv_("codigos_cambio_significado_revisado.csv"),
    "Estudios": csv_("estudios.csv"),
    "Lideres": csv_("lideres.csv"),
}
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    for name, df in sheets.items():
        df.to_excel(xw, sheet_name=name, index=False)

wb = load_workbook(OUT)
base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
fills = {"C": PatternFill("solid", fgColor="F4CCCC"), "B": PatternFill("solid", fgColor="FCE8B2"),
         "AUSENTE": PatternFill("solid", fgColor="D9D9D9")}
edit_fill = PatternFill("solid", fgColor="FFFF00")
widths = {"Leyenda": {"A": 30, "B": 120}}

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = base
            c.alignment = Alignment(vertical="top", wrap_text=ws.title in ("Leyenda",))
    for c in ws[1]:
        c.font = bold
    ws.freeze_panes = "A2"
    if ws.title != "Leyenda":
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    for col, w in widths.get(ws.title, {}).items():
        ws.column_dimensions[col].width = w

ws = wb["Crosswalk"]
hdr = {c.value: c.column for c in ws[1]}
for r in range(2, ws.max_row + 1):
    lvl = ws.cell(r, hdr["nivel_final"]).value
    if lvl in fills:
        ws.cell(r, hdr["nivel_final"]).fill = fills[lvl]
    ws.cell(r, hdr["comentario"]).fill = edit_fill
for col, w in (("etiqueta", 40), ("cambios", 60), ("nota", 90), ("comentario", 40)):
    ws.column_dimensions[get_column_letter(hdr[col])].width = w

for name, col, w in (("Enunciados", "enunciado", 120), ("Rupturas", "descripcion", 90), ("Decisiones", "nota", 120),
                     ("Partidos", "variantes", 50), ("Partidos", "nota", 70), ("Problemas", "etiqueta", 60),
                     ("Problemas", "nota", 60), ("Codigos", "variantes", 90), ("Codigos", "resolucion", 70),
                     ("Codigos", "verificacion", 60)):
    ws = wb[name]
    hdr = {c.value: c.column for c in ws[1]}
    if col in hdr:
        ws.column_dimensions[get_column_letter(hdr[col])].width = w
        for r in range(2, ws.max_row + 1):
            ws.cell(r, hdr[col]).alignment = Alignment(wrap_text=True, vertical="top")

ws = wb["Rupturas"]
hdr = {c.value: c.column for c in ws[1]}
for r in range(2, ws.max_row + 1):
    t = ws.cell(r, hdr["tipo"]).value
    if t in fills:
        ws.cell(r, hdr["tipo"]).fill = fills[t]

wb.save(OUT)
print("ok", OUT)
